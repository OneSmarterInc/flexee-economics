from __future__ import annotations

import csv
import hashlib
import json
import os
import shutil
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
PACKAGE = ROOT / "halden-week4-data-package"
SOURCE_ROOT = Path(os.environ["HALDEN_SOURCE_ROOT"]) if os.environ.get("HALDEN_SOURCE_ROOT") else None
EXTERNAL_SOURCE_HASHES = {
    "halden-constants-ledger.md": "62D26482535FCD6CA8D9EDC5644D8784CBDCF9F3879B6CC87A30392251328BBC",
    "halden-week4-data-package-spec.md": "8BA320474AAB95A13B72CCF854A335AC005DB42A799121B0B1957B51CD01CE55",
}


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest().upper()


def read_params(path: Path) -> dict[str, Decimal]:
    with path.open(newline="", encoding="utf-8") as f:
        return {row["parameter"]: Decimal(row["value"]) for row in csv.DictReader(f)}


def read_lift_2024(path: Path) -> Decimal:
    with path.open(newline="", encoding="utf-8") as f:
        for row in csv.DictReader(f):
            if row["vintage"] == "2024":
                return Decimal(row["cash_lifting_cost_usd_bbl"])
    raise RuntimeError("2024 lifting-cost row missing")


def q2(value: Decimal) -> str:
    return str(value.quantize(Decimal("0.00")))


def build_outputs() -> tuple[dict, dict]:
    const = read_params(PACKAGE / "data" / "cost_constants.csv")
    prior = read_params(PACKAGE / "data" / "worked_example_prior.csv")
    lift_2024 = read_lift_2024(PACKAGE / "data" / "permian_lifting.csv")

    worked_realized = prior["wti_prior"] - prior["permian_wellhead_discount"]
    worked_product = worked_realized + prior["gc_crack_321"] + prior["br_complexity_premium"]
    worked_integrated = worked_product - prior["delivered_marginal_cost"] - prior["br_opex"]
    worked_market = worked_realized + prior["transport_permian_to_br"]
    worked_marginal = prior["delivered_marginal_cost"] + prior["sr_capital_charge"]

    current_mc = lift_2024 + const["gathering_cost"] + const["transport_permian_to_br"]
    current_realized = const["wti"] - const["permian_wellhead_discount"]
    current_product = current_realized + const["gc_crack_321"] + const["br_complexity_premium"]
    current_integrated = current_product - current_mc - const["br_opex"]
    current_market = current_realized + const["transport_permian_to_br"]
    current_marginal = current_mc + const["sr_capital_charge"]
    current_midpoint = (current_market + current_marginal) / Decimal("2")
    capture = (current_market - current_midpoint) * const["geneva_capture_rate"]
    daily_capture = capture * const["geneva_max_volume"]

    up_target = Decimal("45.00")
    ref_target = Decimal("30.00")

    def split(tp: Decimal) -> dict[str, str]:
        upstream = tp - current_mc
        refining = current_product - tp - const["br_opex"]
        return {
            "transfer_price": q2(tp),
            "upstream_margin": q2(upstream),
            "refining_margin": q2(refining),
            "upstream_vs_target": q2(upstream - up_target),
            "refining_vs_target": q2(refining - ref_target),
            "integrated_margin": q2(upstream + refining),
        }

    worked = {
        "metadata": {
            "name": "Week 4 worked example",
            "source": "worked_example_prior.csv and halden_week4.xlsx",
            "units": "usd_bbl unless otherwise noted",
        },
        "inputs": {
            "wti": q2(prior["wti_prior"]),
            "permian_wellhead_discount": q2(prior["permian_wellhead_discount"]),
            "transport_permian_to_br": q2(prior["transport_permian_to_br"]),
            "delivered_marginal_cost": q2(prior["delivered_marginal_cost"]),
            "sr_capital_charge": q2(prior["sr_capital_charge"]),
            "gc_crack_321": q2(prior["gc_crack_321"]),
            "br_complexity_premium": q2(prior["br_complexity_premium"]),
            "br_opex": q2(prior["br_opex"]),
        },
        "outputs": {
            "realized_wellhead": q2(worked_realized),
            "product_slate_value": q2(worked_product),
            "integrated_margin": q2(worked_integrated),
            "transfer_prices": {
                "market": q2(worked_market),
                "marginal_cost": q2(worked_marginal),
            },
            "segment_splits": {
                "market": {
                    "upstream_margin": q2(worked_market - prior["delivered_marginal_cost"]),
                    "refining_margin": q2(worked_product - worked_market - prior["br_opex"]),
                    "integrated_margin": q2(worked_integrated),
                },
                "marginal_cost": {
                    "upstream_margin": q2(worked_marginal - prior["delivered_marginal_cost"]),
                    "refining_margin": q2(worked_product - worked_marginal - prior["br_opex"]),
                    "integrated_margin": q2(worked_integrated),
                },
            },
        },
    }

    reference = {
        "metadata": {
            "name": "Week 4 current reference outputs",
            "units": "usd_bbl unless otherwise noted",
            "package_version": "2026-09-22-normalized",
            "authority_order": [
                "halden-constants-ledger.md",
                "halden-week4-data-package-spec.md",
                "halden-week4-data-package package files",
            ],
        },
        "inputs": {
            "wti": q2(const["wti"]),
            "permian_wellhead_discount": q2(const["permian_wellhead_discount"]),
            "gathering_cost": q2(const["gathering_cost"]),
            "transport_permian_to_br": q2(const["transport_permian_to_br"]),
            "cash_lifting_cost_2024": q2(lift_2024),
            "sr_capital_charge": q2(const["sr_capital_charge"]),
            "gc_crack_321": q2(const["gc_crack_321"]),
            "br_complexity_premium": q2(const["br_complexity_premium"]),
            "br_opex": q2(const["br_opex"]),
            "geneva_capture_rate": str(const["geneva_capture_rate"]),
            "geneva_max_volume_bbl_day": str(const["geneva_max_volume"]),
        },
        "outputs": {
            "realized_wellhead": q2(current_realized),
            "product_slate_value": q2(current_product),
            "delivered_marginal_cost": q2(current_mc),
            "integrated_margin": q2(current_integrated),
            "transfer_prices": {
                "market": q2(current_market),
                "marginal_cost": q2(current_marginal),
                "lazy_midpoint": q2(current_midpoint),
            },
            "segment_splits": {
                "market": split(current_market),
                "marginal_cost": split(current_marginal),
                "lazy_midpoint": split(current_midpoint),
            },
            "geneva_arbitrage": {
                "market_to_midpoint_gap": q2(current_market - current_midpoint),
                "capture_rate": str(const["geneva_capture_rate"]),
                "capture_per_bbl": str(capture.normalize()),
                "max_volume_bbl_day": int(const["geneva_max_volume"]),
                "daily_capture_at_volume_cap": str(daily_capture.quantize(Decimal("0.001"))),
                "time_basis": "unresolved; source only supplies bbl/day volume cap",
            },
        },
        "unresolved": [
            "custom transfer-price minimum, maximum, increment, and precision",
            "rounding method for 9.625/bbl Geneva capture display",
            "Geneva period/day-count convention beyond bbl/day cap",
            "Week 4 to Week 6 cohort classification thresholds",
        ],
    }
    return worked, reference


def write_json(path: Path, data: dict) -> None:
    path.write_text(json.dumps(data, indent=4) + "\n", encoding="utf-8")


def notebook(cells: list[dict]) -> dict:
    return {
        "cells": cells,
        "metadata": {
            "language_info": {"name": "python", "pygments_lexer": "ipython3"},
            "kernelspec": {"display_name": "Python 3", "language": "python", "name": "python3"},
        },
        "nbformat": 4,
        "nbformat_minor": 5,
    }


def md(source: str) -> dict:
    return {"cell_type": "markdown", "metadata": {}, "source": source.splitlines(keepends=True)}


def code(source: str) -> dict:
    return {
        "cell_type": "code",
        "execution_count": None,
        "metadata": {},
        "outputs": [],
        "source": source.splitlines(keepends=True),
    }


def write_student_notebook() -> None:
    cells = [
        md("""# Halden Energy - Week 4: Transfer Pricing
**Student analysis notebook.**

Run this notebook from the package root (`halden-week4-data-package/`) or from the `student/` directory. It loads the canonical CSVs in `data/`.

The current-week analysis cells intentionally remain TODOs for students."""),
        code("""from pathlib import Path
import pandas as pd

PACKAGE_ROOT = Path.cwd()
if not (PACKAGE_ROOT / "data").exists() and (PACKAGE_ROOT.parent / "data").exists():
    PACKAGE_ROOT = PACKAGE_ROOT.parent
DATA_DIR = PACKAGE_ROOT / "data"

lift = pd.read_csv(DATA_DIR / "permian_lifting.csv")
const = pd.read_csv(DATA_DIR / "cost_constants.csv")
comp = pd.read_csv(DATA_DIR / "segment_comp.csv")

def cv(name, df=const):
    return float(df.loc[df["parameter"] == name, "value"].iloc[0])

lift"""),
        md("""## Part 1 - Data

`permian_lifting.csv` gives lifting cost by vintage. `cost_constants.csv` combines several conceptual Week 4 datasets from the spec. `segment_comp.csv` gives target margins."""),
        code("const"),
        code("comp"),
        md("""## Part 2 - Worked example

The prior-period worked example is fully solved so you can see the method before applying it to the current week."""),
        code("""we = pd.read_csv(DATA_DIR / "worked_example_prior.csv")

def wev(name):
    return float(we.loc[we["parameter"] == name, "value"].iloc[0])

wti_p = wev("wti_prior")
disc = wev("permian_wellhead_discount")
transport = wev("transport_permian_to_br")
mc = wev("delivered_marginal_cost")
srcap = wev("sr_capital_charge")
crack = wev("gc_crack_321")
compl = wev("br_complexity_premium")
opex = wev("br_opex")

realized = wti_p - disc
product_val = realized + crack + compl
integrated = product_val - mc - opex

print(f"Realized wellhead:   ${realized:.2f}")
print(f"Product slate value: ${product_val:.2f}")
print(f"INTEGRATED margin:   ${integrated:.2f}")"""),
        code("""tp_market = realized + transport
tp_marginal = mc + srcap

def split(tp):
    upstream = tp - mc
    refining = product_val - tp - opex
    return upstream, refining, upstream + refining

for name, tp in [("Market", tp_market), ("Marginal", tp_marginal)]:
    u, r, s = split(tp)
    print(f"{name:9s} TP=${tp:6.2f} | upstream ${u:6.2f} | refining ${r:6.2f} | sum ${s:6.2f}")"""),
        md("""## Part 3 - Your analysis

Fill in the TODO cells using this week's data. Do not rely on the app to do this analysis for you."""),
        code("""wti = cv("wti")
disc = cv("permian_wellhead_discount")
transport = cv("transport_permian_to_br")
gathering = cv("gathering_cost")
srcap = cv("sr_capital_charge")
crack = cv("gc_crack_321")
compl = cv("br_complexity_premium")
opex = cv("br_opex")

lift_2024 = float(lift.loc[lift["vintage"] == 2024, "cash_lifting_cost_usd_bbl"].iloc[0])
mc = lift_2024 + gathering + transport
print(f"Delivered marginal cost (2024 vintage): ${mc:.2f}")"""),
        code("""# TODO: compute the realized wellhead price (WTI minus the wellhead discount)
realized = None

# TODO: compute the product slate value (realized + crack + complexity premium)
product_val = None

# TODO: compute the integrated margin (product value - marginal cost - opex)
integrated = None

# TODO: compute the three transfer price options
tp_market = None
tp_marginal = None
tp_lazy = None"""),
        code("""up_target = float(comp.loc[comp["segment"] == "upstream", "target_margin_usd_bbl"].iloc[0])
ref_target = float(comp.loc[comp["segment"] == "refining", "target_margin_usd_bbl"].iloc[0])

def analyse(tp):
    upstream = None
    refining = None
    up_vs = None
    ref_vs = None
    return upstream, refining, up_vs, ref_vs"""),
        code("""capture_rate = cv("geneva_capture_rate")
# TODO: arb_per_bbl = capture_rate * (tp_market - tp_lazy)"""),
        md("""## Your decision

Use the analysis to write the memo: key assumptions, analytical method and result, and decision logic."""),
    ]
    write_json(PACKAGE / "student" / "halden_week4_analysis.ipynb", notebook(cells))


def write_faculty_notebook() -> None:
    cells = [
        md("""# Halden Energy - Week 4 Faculty Solution

This notebook loads the same canonical CSV files as the student notebook and completes the current-week Week 4 analysis. It does not introduce economics beyond the constants ledger, Week 4 spec, and supplied package files."""),
        code("""from pathlib import Path
from decimal import Decimal
import pandas as pd

PACKAGE_ROOT = Path.cwd()
if not (PACKAGE_ROOT / "data").exists() and (PACKAGE_ROOT.parent / "data").exists():
    PACKAGE_ROOT = PACKAGE_ROOT.parent
DATA_DIR = PACKAGE_ROOT / "data"

lift = pd.read_csv(DATA_DIR / "permian_lifting.csv", dtype=str)
const = pd.read_csv(DATA_DIR / "cost_constants.csv", dtype=str)
comp = pd.read_csv(DATA_DIR / "segment_comp.csv", dtype=str)

def cv(name):
    return Decimal(const.loc[const["parameter"] == name, "value"].iloc[0])

def money(value):
    return value.quantize(Decimal("0.00"))

def out(value):
    return str(money(value))"""),
        code("""wti = cv("wti")
disc = cv("permian_wellhead_discount")
transport = cv("transport_permian_to_br")
gathering = cv("gathering_cost")
srcap = cv("sr_capital_charge")
crack = cv("gc_crack_321")
complexity = cv("br_complexity_premium")
opex = cv("br_opex")
capture_rate = cv("geneva_capture_rate")
max_volume = cv("geneva_max_volume")

lift_2024 = Decimal(lift.loc[lift["vintage"] == "2024", "cash_lifting_cost_usd_bbl"].iloc[0])
up_target = Decimal(comp.loc[comp["segment"] == "upstream", "target_margin_usd_bbl"].iloc[0])
ref_target = Decimal(comp.loc[comp["segment"] == "refining", "target_margin_usd_bbl"].iloc[0])

delivered_marginal_cost = lift_2024 + gathering + transport
realized_wellhead = wti - disc
product_slate_value = realized_wellhead + crack + complexity
integrated_margin = product_slate_value - delivered_marginal_cost - opex

tp_market = realized_wellhead + transport
tp_marginal = delivered_marginal_cost + srcap
tp_midpoint = (tp_market + tp_marginal) / Decimal("2")

def split(tp):
    upstream = tp - delivered_marginal_cost
    refining = product_slate_value - tp - opex
    return {
        "transfer_price": out(tp),
        "upstream_margin": out(upstream),
        "refining_margin": out(refining),
        "upstream_vs_target": out(upstream - up_target),
        "refining_vs_target": out(refining - ref_target),
        "integrated_margin": out(upstream + refining),
    }

geneva_gap = tp_market - tp_midpoint
geneva_capture_per_bbl = geneva_gap * capture_rate

reference_outputs = {
    "delivered_marginal_cost": out(delivered_marginal_cost),
    "transfer_prices": {
        "market": out(tp_market),
        "marginal_cost": out(tp_marginal),
        "lazy_midpoint": out(tp_midpoint),
    },
    "integrated_margin": out(integrated_margin),
    "segment_splits": {
        "market": split(tp_market),
        "marginal_cost": split(tp_marginal),
        "lazy_midpoint": split(tp_midpoint),
    },
    "geneva_arbitrage": {
        "gap": out(geneva_gap),
        "capture_rate": str(capture_rate),
        "capture_per_bbl": str(geneva_capture_per_bbl.normalize()),
        "max_volume_bbl_day": int(max_volume),
        "time_basis": "unresolved; do not annualize or monthly-convert without an authoritative rule",
    },
}

reference_outputs"""),
        code("""assert reference_outputs["delivered_marginal_cost"] == "14.10"
assert reference_outputs["transfer_prices"] == {
    "market": "73.70",
    "marginal_cost": "18.70",
    "lazy_midpoint": "46.20",
}
assert reference_outputs["integrated_margin"] == "76.75"
assert reference_outputs["segment_splits"]["market"]["upstream_margin"] == "59.60"
assert reference_outputs["segment_splits"]["market"]["refining_margin"] == "17.15"
assert reference_outputs["segment_splits"]["marginal_cost"]["upstream_margin"] == "4.60"
assert reference_outputs["segment_splits"]["marginal_cost"]["refining_margin"] == "72.15"
assert reference_outputs["segment_splits"]["lazy_midpoint"]["upstream_margin"] == "32.10"
assert reference_outputs["segment_splits"]["lazy_midpoint"]["refining_margin"] == "44.65"
assert reference_outputs["geneva_arbitrage"]["gap"] == "27.50"
assert reference_outputs["geneva_arbitrage"]["capture_per_bbl"] == "9.625"
print("Week 4 faculty solution checks passed")"""),
    ]
    write_json(PACKAGE / "faculty" / "halden_week4_solution.ipynb", notebook(cells))


def write_faculty_workbook() -> None:
    src = PACKAGE / "student" / "halden_week4.xlsx"
    dst = PACKAGE / "faculty" / "halden_week4_solution.xlsx"
    shutil.copyfile(src, dst)
    wb = load_workbook(dst)
    ws = wb["Your Analysis"]
    formulas = {
        "B15": "=B5-B6",
        "B16": "=B15+B10+B11",
        "B17": "=B16-B8-B12",
        "B18": "=B15+B7",
        "B19": "=B8+B9",
        "B20": "=(B18+B19)/2",
        "B24": "=B18-B8",
        "C24": "=B16-B18-B12",
        "D24": "=B24-'Data — Compensation'!C5",
        "E24": "=C24-'Data — Compensation'!C6",
        "B25": "=B19-B8",
        "C25": "=B16-B19-B12",
        "D25": "=B25-'Data — Compensation'!C5",
        "E25": "=C25-'Data — Compensation'!C6",
        "B26": "=B20-B8",
        "C26": "=B16-B20-B12",
        "D26": "=B26-'Data — Compensation'!C5",
        "E26": "=C26-'Data — Compensation'!C6",
        "B30": "='Data — Permian'!B18*(B18-B20)",
    }
    for cell, formula in formulas.items():
        ws[cell] = formula
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.save(dst)


def write_manifest_and_readme() -> None:
    manifest = """# Halden Energy - Week 4 Data Package Manifest

## Purpose

Week 4 transfer-pricing reference package for the Halden Managerial Economics simulation. This package normalizes the supplied flat Week 4 materials into one package root without changing authoritative economic values.

## Authority Order

1. `halden-constants-ledger.md`
2. `halden-week4-data-package-spec.md`
3. this package's files

If a package file disagrees with the constants ledger, flag the conflict. Do not silently edit economic values.

## Canonical Sources

The canonical CSVs live in `data/`.

| File | Conceptual Week 4 datasets represented |
| --- | --- |
| `data/permian_lifting.csv` | Permian lifting costs by vintage |
| `data/cost_constants.csv` | Baton Rouge yield economics, Gulf Coast market differentials, Geneva arbitrage capability, and cost constants |
| `data/segment_comp.csv` | Segment compensation plan |
| `data/worked_example_prior.csv` | Integrated-margin worked-example basis |

The Week 4 spec describes six conceptual datasets. This package intentionally preserves the supplied four-CSV layout instead of splitting data solely to match that conceptual count.

## Student Deliverables

- `student/halden_week4.xlsx`
- `student/halden_week4_analysis.ipynb`

The student workbook embeds CSV-derived values in sheets and does not dynamically read the CSV files. The student notebook loads the canonical CSVs from `data/`.

## Faculty Deliverables

- `faculty/halden_week4_solution.xlsx`
- `faculty/halden_week4_solution.ipynb`

Both faculty deliverables use the same source data and complete the current-week analysis.

## Expected Outputs

- `expected/worked_example.json`
- `expected/week4_reference.json`

The prior-period worked example must reproduce `$68.25` integrated margin.

Current Week 4 invariants:

- delivered marginal cost: `14.10`
- marginal transfer price: `18.70`
- market transfer price: `73.70`
- midpoint transfer price: `46.20`
- integrated margin: `76.75`
- Geneva midpoint capture: `9.625/bbl`

## Unresolved Rules

The package does not define custom transfer-price min/max/increment/precision.

The package does not define a monthly or annual Geneva time basis. Keep Geneva capture as `9.625 x applicable barrels` until an authoritative period rule is supplied.

The package does not define Week 4 to Week 6 cohort classification thresholds.

## Validation

From the application repository root:

```bash
python scripts/validate_week4_package.py
```

The validation script checks required files, source hashes, CSV arithmetic, worked-example output, student/faculty notebook execution, expected-output consistency, and workbook formulas.
"""
    readme = """# Halden Week 4 Data Package

This normalized package repairs the supplied Week 4 materials into a self-contained package root.

Run notebooks from this package root or from their own `student/` / `faculty/` directories. Both notebooks resolve `data/` by relative path and do not use machine-specific absolute paths.

Install the minimal Python dependency with:

```bash
pip install -r requirements.txt
```

Then run:

```bash
python ../scripts/validate_week4_package.py
```
"""
    (PACKAGE / "MANIFEST.md").write_text(manifest, encoding="utf-8")
    (PACKAGE / "README.md").write_text(readme, encoding="utf-8")
    (PACKAGE / "requirements.txt").write_text("pandas\nopenpyxl\n", encoding="utf-8")


def write_expected_outputs(worked: dict, reference: dict) -> None:
    external_hashes = dict(EXTERNAL_SOURCE_HASHES)
    if SOURCE_ROOT is not None:
        external_hashes = {
            name: sha256(SOURCE_ROOT / name)
            for name in EXTERNAL_SOURCE_HASHES
        }
    hashes = {
        **external_hashes,
        "data/permian_lifting.csv": sha256(PACKAGE / "data" / "permian_lifting.csv"),
        "data/cost_constants.csv": sha256(PACKAGE / "data" / "cost_constants.csv"),
        "data/segment_comp.csv": sha256(PACKAGE / "data" / "segment_comp.csv"),
        "data/worked_example_prior.csv": sha256(PACKAGE / "data" / "worked_example_prior.csv"),
        "student/halden_week4.xlsx": sha256(PACKAGE / "student" / "halden_week4.xlsx"),
        "student/halden_week4_analysis.ipynb": sha256(PACKAGE / "student" / "halden_week4_analysis.ipynb"),
        "faculty/halden_week4_solution.ipynb": sha256(PACKAGE / "faculty" / "halden_week4_solution.ipynb"),
        "faculty/halden_week4_solution.xlsx": sha256(PACKAGE / "faculty" / "halden_week4_solution.xlsx"),
    }
    for doc in (worked, reference):
        doc["provenance"] = {
            "source_hashes": hashes,
            "package_root": "halden-week4-data-package",
        }
    write_json(PACKAGE / "expected" / "worked_example.json", worked)
    write_json(PACKAGE / "expected" / "week4_reference.json", reference)


def main() -> None:
    worked, reference = build_outputs()
    write_student_notebook()
    write_faculty_notebook()
    write_faculty_workbook()
    write_manifest_and_readme()
    worked, reference = build_outputs()
    write_expected_outputs(worked, reference)


if __name__ == "__main__":
    main()
