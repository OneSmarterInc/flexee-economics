from __future__ import annotations

import csv
import hashlib
import json
import os
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
PACKAGE = ROOT / "halden-week8-data-package"


def fail(message: str) -> None:
    raise AssertionError(message)


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)

    return h.hexdigest().lower()


def load_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def load_csv(path: Path) -> list[dict[str, str]]:
    with path.open(newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def execute_notebook(path: Path, cwd: Path) -> dict:
    nb = load_json(path)
    namespace: dict = {"__name__": "__notebook__"}
    old_cwd = Path.cwd()
    try:
        os.chdir(cwd)
        for index, cell in enumerate(nb["cells"]):
            if cell.get("cell_type") != "code":
                continue
            source = "".join(cell.get("source", []))
            exec(compile(source, f"{path.name}:cell-{index}", "exec"), namespace)
    finally:
        os.chdir(old_cwd)

    return namespace


def assert_close(actual: float, expected: float, precision: int = 2) -> None:
    if round(actual, precision) != round(expected, precision):
        fail(f"Expected {expected}, got {actual}")


def validate_required_files() -> None:
    required = [
        "MANIFEST.md",
        "halden_week8.xlsx",
        "halden_week8_analysis.ipynb",
        "data/baseline_state.csv",
        "data/compliance_history.csv",
        "data/opec_scenarios.csv",
        "data/propagation_coefficients.csv",
        "data/worked_example_prior.csv",
        "faculty/halden_week8_FACULTY_SOLUTION.xlsx",
        "fixtures/provenance.json",
        "fixtures/week8_golden.json",
    ]
    missing = [path for path in required if not (PACKAGE / path).exists()]
    if missing:
        fail(f"Missing package files: {missing}")


def validate_hashes() -> None:
    provenance = load_json(PACKAGE / "fixtures" / "provenance.json")
    for rel_path, expected_hash in provenance["artifacts"].items():
        path = PACKAGE / rel_path
        if not path.exists():
            fail(f"Provenance references missing artifact: {rel_path}")

        actual_hash = sha256(path)
        if actual_hash != expected_hash.lower():
            fail(f"Hash mismatch for {rel_path}: {actual_hash} != {expected_hash}")


def validate_csv_schemas() -> None:
    expected = {
        "baseline_state.csv": ["parameter", "value", "unit", "note"],
        "compliance_history.csv": ["episode", "announced_cut_mbd", "delivered_mbd", "compliance_pct", "note"],
        "opec_scenarios.csv": ["scenario", "probability", "wti_resolved", "delta_wti"],
        "propagation_coefficients.csv": ["channel", "coefficient", "unit", "note"],
        "worked_example_prior.csv": ["scenario", "probability", "delta_wti"],
    }

    for filename, columns in expected.items():
        rows = load_csv(PACKAGE / "data" / filename)
        if not rows:
            fail(f"{filename} has no data rows.")
        if list(rows[0].keys()) != columns:
            fail(f"{filename} columns mismatch: {list(rows[0].keys())}")

        unique_rows = {tuple(row.items()) for row in rows}
        if len(unique_rows) != len(rows):
            fail(f"{filename} contains duplicate rows.")

        for row in rows:
            missing = [key for key, value in row.items() if value == ""]
            if missing:
                fail(f"{filename} row has missing values in {missing}: {row}")


def validate_golden_math() -> None:
    golden = load_json(PACKAGE / "fixtures" / "week8_golden.json")
    baseline = {row["parameter"]: float(row["value"]) for row in load_csv(PACKAGE / "data" / "baseline_state.csv")}
    coefficients = {row["channel"]: float(row["coefficient"]) for row in load_csv(PACKAGE / "data" / "propagation_coefficients.csv")}
    scenarios = load_csv(PACKAGE / "data" / "opec_scenarios.csv")

    for key, expected in golden["propagation_coefficients"].items():
        assert_close(coefficients[key], expected, precision=4)

    expected_wti = sum(float(row["probability"]) * float(row["wti_resolved"]) for row in scenarios)
    assert_close(expected_wti, golden["expected_wti"], precision=2)

    probability_sum = sum(float(row["probability"]) for row in scenarios)
    assert_close(probability_sum, 1.0, precision=6)

    for row in scenarios:
        scenario = row["scenario"]
        delta_wti = float(row["delta_wti"])
        fixture = golden["per_scenario"][scenario]
        upstream = coefficients["upstream_realization"] * delta_wti
        crack = baseline["crack_base"] + coefficients["refining_crack"] * delta_wti
        retail = (
            coefficients["retail_demand_elasticity"]
            * (coefficients["retail_passthrough"] * delta_wti / 42.0 / baseline["pump_base"])
            * 100
        )

        assert_close(float(row["probability"]), fixture["sim_probability"], precision=6)
        assert_close(float(row["wti_resolved"]), baseline["wti_pre"] + delta_wti, precision=2)
        assert_close(upstream, fixture["upstream_per_bbl"], precision=2)
        assert_close(crack, fixture["crack"], precision=2)
        assert_close(retail, fixture["retail_vol_pct"], precision=3)

    if not all(assertion["holds"] for assertion in golden["ordering_assertions"]):
        fail("One or more golden ordering assertions is false.")


def validate_notebook() -> None:
    namespace = execute_notebook(PACKAGE / "halden_week8_analysis.ipynb", PACKAGE)
    for symbol in ["pd", "sc", "co", "b", "ch", "propagate", "ev", "exp_upstream", "exp_crack"]:
        if symbol not in namespace:
            fail(f"Student notebook did not define {symbol}.")

    up, crack, retail = namespace["propagate"](14)
    assert_close(up, 14.0, precision=2)
    assert_close(crack, 16.6, precision=2)
    assert_close(retail, -0.312, precision=3)
    assert_close(namespace["ev"], 80.7, precision=2)


def validate_workbooks() -> None:
    expected_sheets = [
        "README",
        "Data \u2014 Scenarios & Propagation",
        "Data \u2014 Compliance",
        "Worked Example",
        "Your Analysis",
    ]
    student = load_workbook(PACKAGE / "halden_week8.xlsx", data_only=False)
    student_values = load_workbook(PACKAGE / "halden_week8.xlsx", data_only=True)
    faculty = load_workbook(PACKAGE / "faculty" / "halden_week8_FACULTY_SOLUTION.xlsx", data_only=False)
    faculty_values = load_workbook(PACKAGE / "faculty" / "halden_week8_FACULTY_SOLUTION.xlsx", data_only=True)

    if student.sheetnames != expected_sheets:
        fail(f"Student workbook sheets mismatch: {student.sheetnames}")
    if faculty.sheetnames != expected_sheets:
        fail(f"Faculty workbook sheets mismatch: {faculty.sheetnames}")
    if any(sheet.sheet_state != "visible" for sheet in student.worksheets):
        fail("Student workbook contains hidden sheets.")
    if any(sheet.sheet_state != "visible" for sheet in faculty.worksheets):
        fail("Faculty workbook contains hidden sheets.")
    if len(getattr(student, "_external_links", [])) != 0:
        fail("Student workbook contains external links.")
    if len(getattr(faculty, "_external_links", [])) != 0:
        fail("Faculty workbook contains external links.")

    student_analysis = student_values["Your Analysis"]
    if any(student_analysis[cell].value is not None for cell in ["B6", "B7", "B8"]):
        fail("Student workbook exposes solved current-week probabilities.")

    faculty_analysis = faculty_values["Your Analysis"]
    assert_close(faculty_analysis["B6"].value, 0.35, precision=6)
    assert_close(faculty_analysis["B7"].value, 0.40, precision=6)
    assert_close(faculty_analysis["B8"].value, 0.25, precision=6)
    assert_close(faculty_analysis["B9"].value, 1.0, precision=6)

    for workbook in [student, faculty]:
        for worksheet in workbook.worksheets:
            for row in worksheet.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str) and cell.value.startswith("#"):
                        fail(f"Workbook formula error marker in {worksheet.title}!{cell.coordinate}: {cell.value}")

    formulas = faculty["Worked Example"]
    expected_formulas = {
        "D5": "=C5*'Data \u2014 Scenarios & Propagation'!B11",
        "E5": "='Data \u2014 Scenarios & Propagation'!B19+C5*'Data \u2014 Scenarios & Propagation'!B12",
        "F8": "=SUM(F5:F7)",
    }
    for cell, formula in expected_formulas.items():
        if formulas[cell].value != formula:
            fail(f"Faculty workbook {cell} formula mismatch.")


def main() -> None:
    validate_required_files()
    validate_hashes()
    validate_csv_schemas()
    validate_golden_math()
    validate_notebook()
    validate_workbooks()
    print("Week 8 package validation passed")


if __name__ == "__main__":
    main()
