from __future__ import annotations

import csv
import hashlib
import json
import os
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
PACKAGE = ROOT / "halden-week6-data-package"


def fail(message: str) -> None:
    raise AssertionError(message)


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)

    return h.hexdigest().lower()


def load_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def load_csv(path: Path) -> list[dict[str, str]]:
    with path.open(newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def npv(rate: float, flows: list[float]) -> float:
    return sum(flow / ((1 + rate) ** year) for year, flow in enumerate(flows))


def irr(flows: list[float]) -> float:
    lo, hi = -0.9, 3.0
    for _ in range(300):
        mid = (lo + hi) / 2
        if npv(mid, flows) > 0:
            lo = mid
        else:
            hi = mid

    return mid


def execute_notebook(path: Path, cwd: Path) -> dict:
    nb = load_json(path)
    namespace: dict = {"__name__": "__notebook__"}
    old_cwd = Path.cwd()
    try:
        os.chdir(cwd)
        for index, cell in enumerate(nb["cells"]):
            if cell.get("cell_type") != "code":
                continue
            source = "".join(cell.get("source", []))
            exec(compile(source, f"{path.name}:cell-{index}", "exec"), namespace)
    finally:
        os.chdir(old_cwd)

    return namespace


def assert_close(actual: float, expected: float, precision: int = 2) -> None:
    if round(actual, precision) != round(expected, precision):
        fail(f"Expected {expected}, got {actual}")


def validate_required_files() -> None:
    required = [
        "MANIFEST.md",
        "halden_week6.xlsx",
        "halden_week6_analysis.ipynb",
        "data/cohort_discount_schedule.csv",
        "data/cost_of_capital.csv",
        "data/currency_helix.csv",
        "data/forecast_haircuts.csv",
        "data/project_cashflows.csv",
        "data/worked_example_prior.csv",
        "faculty/halden_week6_FACULTY_SOLUTION.xlsx",
        "fixtures/provenance.json",
        "fixtures/week6_golden.json",
    ]
    missing = [path for path in required if not (PACKAGE / path).exists()]
    if missing:
        fail(f"Missing package files: {missing}")


def validate_hashes() -> None:
    provenance = load_json(PACKAGE / "fixtures" / "provenance.json")
    for rel_path, expected_hash in provenance["artifacts"].items():
        path = PACKAGE / rel_path
        if not path.exists():
            fail(f"Provenance references missing artifact: {rel_path}")

        actual_hash = sha256(path)
        if actual_hash != expected_hash.lower():
            fail(f"Hash mismatch for {rel_path}: {actual_hash} != {expected_hash}")


def validate_fixture_math() -> None:
    golden = load_json(PACKAGE / "fixtures" / "week6_golden.json")
    rows = load_csv(PACKAGE / "data" / "project_cashflows.csv")
    projects = {
        "baton_rouge": [float(row["baton_rouge"]) for row in rows],
        "rotterdam": [float(row["rotterdam"]) for row in rows],
        "helix": [float(row["helix"]) for row in rows],
    }

    for project, expected in golden["irr"].items():
        assert_close(irr(projects[project]), expected, precision=4)

    schedule = load_csv(PACKAGE / "data" / "cohort_discount_schedule.csv")
    for row in schedule:
        cohort = row["cohort_wk4_behavior"]
        rate = float(row["discount_rate"])
        fixture = golden["npv_by_cohort"][cohort]
        assert_close(rate, fixture["rate"], precision=4)
        assert_close(float(row["capital_envelope_musd"]), fixture["envelope_musd"], precision=2)

        calculated = {project: npv(rate, flows) for project, flows in projects.items()}
        ranking = sorted(calculated, key=lambda key: -calculated[key])
        if ranking != golden["ranking_by_cohort"][cohort]:
            fail(f"{cohort} ranking mismatch: {ranking}")

        for project, expected in fixture["npv"].items():
            assert_close(calculated[project], expected, precision=2)

    if not all(assertion["holds"] for assertion in golden["ordering_assertions"]):
        fail("One or more golden ordering assertions is false.")


def validate_notebook() -> None:
    namespace = execute_notebook(PACKAGE / "halden_week6_analysis.ipynb", PACKAGE)
    for symbol in ["cf", "sched", "haircuts", "projects", "npv", "irr"]:
        if symbol not in namespace:
            fail(f"Student notebook did not define {symbol}.")

    helix = namespace["projects"]["helix"]
    if round(namespace["irr"](helix), 4) != 0.1240:
        fail("Student notebook IRR helper does not reproduce the Helix fixture.")


def validate_workbooks() -> None:
    expected_sheets = ["README", "Data \u2014 Projects", "Data \u2014 Adjustments", "Worked Example", "Your Analysis"]
    student = load_workbook(PACKAGE / "halden_week6.xlsx", data_only=False)
    faculty_formula = load_workbook(PACKAGE / "faculty" / "halden_week6_FACULTY_SOLUTION.xlsx", data_only=False)
    faculty_values = load_workbook(PACKAGE / "faculty" / "halden_week6_FACULTY_SOLUTION.xlsx", data_only=True)

    if student.sheetnames != expected_sheets:
        fail(f"Student workbook sheets mismatch: {student.sheetnames}")
    if faculty_formula.sheetnames != expected_sheets:
        fail(f"Faculty workbook sheets mismatch: {faculty_formula.sheetnames}")
    if len(getattr(student, "_external_links", [])) != 0:
        fail("Student workbook contains external links.")
    if len(getattr(faculty_formula, "_external_links", [])) != 0:
        fail("Faculty workbook contains external links.")

    ws_formula = faculty_formula["Your Analysis"]
    required_formulas = {
        "B9": "=NPV($B$5,'Data \u2014 Projects'!B5:B14)+'Data \u2014 Projects'!B4",
        "C9": "=IRR('Data \u2014 Projects'!B4:B14)",
        "B10": "=NPV($B$5,'Data \u2014 Projects'!C5:C14)+'Data \u2014 Projects'!C4",
        "C10": "=IRR('Data \u2014 Projects'!C4:C14)",
        "B11": "=NPV($B$5,'Data \u2014 Projects'!D5:D14)+'Data \u2014 Projects'!D4",
        "C11": "=IRR('Data \u2014 Projects'!D4:D14)",
    }
    for cell, formula in required_formulas.items():
        if ws_formula[cell].value != formula:
            fail(f"Faculty workbook {cell} formula mismatch.")

    golden = load_json(PACKAGE / "fixtures" / "week6_golden.json")
    ws_values = faculty_values["Your Analysis"]
    disciplined = golden["npv_by_cohort"]["disciplined"]["npv"]
    assert_close(ws_values["B9"].value, disciplined["baton_rouge"], precision=2)
    assert_close(ws_values["B10"].value, disciplined["rotterdam"], precision=2)
    assert_close(ws_values["B11"].value, disciplined["helix"], precision=2)
    assert_close(ws_values["C9"].value, golden["irr"]["baton_rouge"], precision=4)
    assert_close(ws_values["C10"].value, golden["irr"]["rotterdam"], precision=4)
    assert_close(ws_values["C11"].value, golden["irr"]["helix"], precision=4)


def main() -> None:
    validate_required_files()
    validate_hashes()
    validate_fixture_math()
    validate_notebook()
    validate_workbooks()
    print("Week 6 package validation passed")


if __name__ == "__main__":
    main()
