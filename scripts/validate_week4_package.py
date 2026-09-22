from __future__ import annotations

import csv
import hashlib
import json
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
PACKAGE = ROOT / "halden-week4-data-package"


def fail(message: str) -> None:
    raise AssertionError(message)


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest().upper()


def load_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def params(path: Path) -> dict[str, Decimal]:
    with path.open(newline="", encoding="utf-8") as f:
        return {row["parameter"]: Decimal(row["value"]) for row in csv.DictReader(f)}


def execute_notebook(path: Path, cwd: Path) -> dict:
    nb = load_json(path)
    namespace: dict = {"__name__": "__notebook__"}
    old_cwd = Path.cwd()
    try:
        import os

        os.chdir(cwd)
        for index, cell in enumerate(nb["cells"]):
            if cell.get("cell_type") != "code":
                continue
            source = "".join(cell.get("source", []))
            exec(compile(source, f"{path.name}:cell-{index}", "exec"), namespace)
    finally:
        import os

        os.chdir(old_cwd)
    return namespace


def q2(value: Decimal) -> Decimal:
    return value.quantize(Decimal("0.00"))


def validate_required_files() -> None:
    required = [
        "MANIFEST.md",
        "README.md",
        "requirements.txt",
        "data/permian_lifting.csv",
        "data/cost_constants.csv",
        "data/segment_comp.csv",
        "data/worked_example_prior.csv",
        "student/halden_week4.xlsx",
        "student/halden_week4_analysis.ipynb",
        "faculty/halden_week4_solution.xlsx",
        "faculty/halden_week4_solution.ipynb",
        "expected/worked_example.json",
        "expected/week4_reference.json",
    ]
    missing = [path for path in required if not (PACKAGE / path).exists()]
    if missing:
        fail(f"Missing package files: {missing}")


def validate_hashes() -> None:
    expected = load_json(PACKAGE / "expected" / "week4_reference.json")
    hashes = expected["provenance"]["source_hashes"]
    for rel, expected_hash in hashes.items():
        path = PACKAGE / rel
        if not path.exists():
            # Ledger/spec hashes are external provenance copied into expected outputs.
            continue
        actual = sha256(path)
        if actual != expected_hash:
            fail(f"Hash mismatch for {rel}: {actual} != {expected_hash}")


def validate_csv_math() -> None:
    const = params(PACKAGE / "data" / "cost_constants.csv")
    prior = params(PACKAGE / "data" / "worked_example_prior.csv")
    reference = load_json(PACKAGE / "expected" / "week4_reference.json")
    worked = load_json(PACKAGE / "expected" / "worked_example.json")

    with (PACKAGE / "data" / "permian_lifting.csv").open(newline="", encoding="utf-8") as f:
        lift_rows = list(csv.DictReader(f))
    lift_2024 = Decimal(next(row["cash_lifting_cost_usd_bbl"] for row in lift_rows if row["vintage"] == "2024"))

    delivered = lift_2024 + const["gathering_cost"] + const["transport_permian_to_br"]
    realized = const["wti"] - const["permian_wellhead_discount"]
    product = realized + const["gc_crack_321"] + const["br_complexity_premium"]
    integrated = product - delivered - const["br_opex"]
    market = realized + const["transport_permian_to_br"]
    marginal = delivered + const["sr_capital_charge"]
    midpoint = (market + marginal) / Decimal("2")
    capture = (market - midpoint) * const["geneva_capture_rate"]

    out = reference["outputs"]
    assert q2(delivered) == Decimal(out["delivered_marginal_cost"])
    assert q2(integrated) == Decimal(out["integrated_margin"])
    assert q2(market) == Decimal(out["transfer_prices"]["market"])
    assert q2(marginal) == Decimal(out["transfer_prices"]["marginal_cost"])
    assert q2(midpoint) == Decimal(out["transfer_prices"]["lazy_midpoint"])
    assert capture == Decimal(out["geneva_arbitrage"]["capture_per_bbl"])

    worked_realized = prior["wti_prior"] - prior["permian_wellhead_discount"]
    worked_product = worked_realized + prior["gc_crack_321"] + prior["br_complexity_premium"]
    worked_integrated = worked_product - prior["delivered_marginal_cost"] - prior["br_opex"]
    assert q2(worked_integrated) == Decimal(worked["outputs"]["integrated_margin"])
    assert q2(worked_integrated) == Decimal("68.25")


def validate_student_notebook() -> None:
    path = PACKAGE / "student" / "halden_week4_analysis.ipynb"
    source_text = path.read_text(encoding="utf-8")
    forbidden_answers = ["76.75", "73.70", "46.20", "59.60", "72.15", "44.65", "9.625"]
    leaked = [answer for answer in forbidden_answers if answer in source_text]
    if leaked:
        fail(f"Student notebook exposes current-week answer values: {leaked}")
    namespace = execute_notebook(path, PACKAGE)
    if namespace.get("integrated") is not None:
        fail("Student notebook current-week integrated value should remain TODO/None.")


def validate_faculty_notebook() -> None:
    namespace = execute_notebook(PACKAGE / "faculty" / "halden_week4_solution.ipynb", PACKAGE)
    actual = namespace.get("reference_outputs")
    if actual is None:
        fail("Faculty notebook did not produce reference_outputs.")
    expected = load_json(PACKAGE / "expected" / "week4_reference.json")["outputs"]
    assert actual["delivered_marginal_cost"] == expected["delivered_marginal_cost"]
    assert actual["transfer_prices"] == expected["transfer_prices"]
    assert actual["integrated_margin"] == expected["integrated_margin"]
    assert actual["segment_splits"] == expected["segment_splits"]
    assert actual["geneva_arbitrage"]["capture_per_bbl"] == expected["geneva_arbitrage"]["capture_per_bbl"]


def validate_workbooks() -> None:
    student = load_workbook(PACKAGE / "student" / "halden_week4.xlsx", data_only=False)
    faculty = load_workbook(PACKAGE / "faculty" / "halden_week4_solution.xlsx", data_only=False)

    expected_sheets = ["README", "Data \u2014 Permian", "Data \u2014 Compensation", "Worked Example", "Your Analysis"]
    assert student.sheetnames == expected_sheets
    assert faculty.sheetnames == expected_sheets
    assert len(getattr(student, "_external_links", [])) == 0
    assert len(getattr(faculty, "_external_links", [])) == 0

    ws = faculty["Your Analysis"]
    expected_formulas = {
        "B15": "=B5-B6",
        "B16": "=B15+B10+B11",
        "B17": "=B16-B8-B12",
        "B18": "=B15+B7",
        "B19": "=B8+B9",
        "B20": "=(B18+B19)/2",
        "B24": "=B18-B8",
        "C24": "=B16-B18-B12",
        "D24": "=B24-'Data \u2014 Compensation'!C5",
        "E24": "=C24-'Data \u2014 Compensation'!C6",
        "B25": "=B19-B8",
        "C25": "=B16-B19-B12",
        "D25": "=B25-'Data \u2014 Compensation'!C5",
        "E25": "=C25-'Data \u2014 Compensation'!C6",
        "B26": "=B20-B8",
        "C26": "=B16-B20-B12",
        "D26": "=B26-'Data \u2014 Compensation'!C5",
        "E26": "=C26-'Data \u2014 Compensation'!C6",
        "B30": "='Data \u2014 Permian'!B18*(B18-B20)",
    }
    for cell, formula in expected_formulas.items():
        assert ws[cell].value == formula, f"{cell} formula mismatch"


def main() -> None:
    validate_required_files()
    validate_hashes()
    validate_csv_math()
    validate_student_notebook()
    validate_faculty_notebook()
    validate_workbooks()
    print("Week 4 package validation passed")


if __name__ == "__main__":
    main()
